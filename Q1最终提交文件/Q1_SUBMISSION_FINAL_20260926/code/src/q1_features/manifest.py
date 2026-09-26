from __future__ import annotations

from dataclasses import asdict, dataclass
from numbers import Integral, Real
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .storage import write_csv, write_json, write_jsonl

REQUIRED_COLUMNS = ("video_id", "clip_id", "text", "label", "annotation")


class ManifestError(ValueError):
    pass


@dataclass(frozen=True)
class Sample:
    sample_index: int
    sample_id: str
    video_id: str
    clip_id: str
    video_relpath: str
    raw_text: str
    excel_sheet: str
    excel_row: int
    input_status: str


def _cell_id(value: Any, *, clip: bool = False) -> tuple[str, bool]:
    if value is None:
        return "", False
    if clip and isinstance(value, Integral):
        return str(int(value)), True
    if clip and isinstance(value, Real) and float(value).is_integer():
        return str(int(value)), True
    return str(value), False


def locate_labels(data_root: str | Path, labels_xlsx: str | Path | None = None) -> Path:
    if labels_xlsx is not None:
        path = Path(labels_xlsx).expanduser().resolve()
        if not path.is_file():
            raise ManifestError(f"labels workbook not found: {path}")
        return path
    root = Path(data_root).expanduser().resolve()
    found = sorted(root.rglob("label-100.xlsx"))
    if not found:
        raise ManifestError(f"no label-100.xlsx below {root}")
    if len(found) == 1:
        return found[0]
    preferred = [p for p in found if p.parent.name == "MOSEI数据集部分原始视频-100条"]
    if len(preferred) == 1:
        return preferred[0]
    raise ManifestError("multiple label-100.xlsx files found; pass --labels-xlsx:\n" + "\n".join(map(str, found)))


def _sheet_table(workbook) -> tuple[Any, int, dict[str, int]]:
    matches: list[tuple[Any, int, dict[str, int]]] = []
    for sheet in workbook.worksheets:
        for row_no, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            names = [str(v).strip() if v is not None else "" for v in row]
            mapping = {name: idx for idx, name in enumerate(names) if name}
            if all(name in mapping for name in REQUIRED_COLUMNS):
                matches.append((sheet, row_no, mapping))
                break
    if len(matches) != 1:
        details = ", ".join(f"{s.title}@{r}" for s, r, _ in matches) or "none"
        raise ManifestError(f"expected exactly one matching worksheet header; found {details}")
    return matches[0]


def _resolve_video(base: Path, video_id: str, clip_id: str, numeric_clip: bool) -> tuple[Path, str]:
    expected = base / video_id / f"{clip_id}.mp4"
    if expected.is_file():
        return expected, "ok"
    parent = expected.parent
    if numeric_clip and parent.is_dir():
        candidates = []
        for p in parent.glob("*.mp4"):
            try:
                if int(p.stem) == int(clip_id):
                    candidates.append(p)
            except ValueError:
                pass
        if len(candidates) == 1:
            return candidates[0], "integer_equivalent"
        if len(candidates) > 1:
            return expected, "ambiguous_path"
    return expected, "missing_media"


def build_manifest(
    data_root: str | Path,
    out_dir: str | Path,
    *,
    labels_xlsx: str | Path | None = None,
) -> list[Sample]:
    labels = locate_labels(data_root, labels_xlsx)
    workbook = load_workbook(labels, read_only=True, data_only=True)
    sheet, header_row, columns = _sheet_table(workbook)
    records: list[dict[str, Any]] = []
    seen: set[str] = set()
    errors: list[str] = []
    for excel_row, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if all(value is None for value in row):
            continue
        video_id, _ = _cell_id(row[columns["video_id"]])
        clip_id, numeric_clip = _cell_id(row[columns["clip_id"]], clip=True)
        raw = row[columns["text"]]
        raw_text = "" if raw is None else str(raw)
        sample_id = f"{video_id}$_${clip_id}"
        status = "ok"
        if not video_id or not clip_id:
            status = "invalid_id"
            errors.append(f"row {excel_row}: empty video_id/clip_id")
        if sample_id in seen:
            status = "duplicate_key"
            errors.append(f"row {excel_row}: duplicate {sample_id}")
        seen.add(sample_id)
        video_path, path_status = _resolve_video(labels.parent, video_id, clip_id, numeric_clip)
        if status == "ok" and path_status != "ok":
            status = path_status
        try:
            relpath = str(video_path.relative_to(Path(data_root).resolve())).replace("\\", "/")
        except ValueError:
            relpath = str(video_path)
        records.append({
            "sample_id": sample_id, "video_id": video_id, "clip_id": clip_id,
            "video_relpath": relpath, "raw_text": raw_text,
            "excel_sheet": sheet.title, "excel_row": excel_row, "input_status": status,
            "label": "" if row[columns["label"]] is None else row[columns["label"]],
            "annotation": "" if row[columns["annotation"]] is None else row[columns["annotation"]],
        })
    records.sort(key=lambda r: (r["video_id"], r["clip_id"]))
    samples = [Sample(sample_index=i, **{k: r[k] for k in Sample.__dataclass_fields__ if k != "sample_index"}) for i, r in enumerate(records)]
    out = Path(out_dir)
    write_jsonl(out / "manifest.jsonl", (asdict(s) for s in samples))
    labels_by_id = {r["sample_id"]: r for r in records}
    write_csv(out / "labels.csv", ({"sample_id": s.sample_id, "label": labels_by_id[s.sample_id]["label"], "annotation": labels_by_id[s.sample_id]["annotation"]} for s in samples), ["sample_id", "label", "annotation"])
    write_csv(out / "sample_index.csv", (asdict(s) for s in samples), list(Sample.__dataclass_fields__))
    sources = len({s.video_id for s in samples})
    write_json(out / "input_report.json", {
        "labels_xlsx": str(labels), "sheet": sheet.title, "sample_count": len(samples),
        "source_count": sources, "expected_sample_count": 100, "expected_source_count": 37,
        "differences": errors + ([] if len(samples) == 100 else [f"sample_count={len(samples)} expected=100"]) + ([] if sources == 37 else [f"source_count={sources} expected=37"]),
    })
    return samples


def load_manifest(path: str | Path) -> list[Sample]:
    from .storage import read_jsonl
    return [Sample(**row) for row in read_jsonl(path)]

